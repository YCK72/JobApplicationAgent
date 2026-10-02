import json
from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from agent import ai, config, db, gmail, policy, reports, server


def test_duplicate_tracking_urls(job):
    same = db.add(dict(job, url=job['url'] + '?utm_source=linkedin'))
    assert same == job['id']
    assert len(db.jobs()) == 1


def test_claim_is_atomic_and_recovery(job):
    assert db.claim(job['id'])
    assert not db.claim(job['id'])
    db.update(job['id'], attempted=1)
    db.recover()
    assert db.get(job['id'])['status'] == 'needs_review'
    assert db.get(job['id'])['attempted'] == 1


@pytest.mark.parametrize('name,expected', [('Microsoft Corporation', True), ('Apple Inc.', True), ('Pineapple Labs', False), ('Alphabet', True)])
def test_major_company(name, expected):
    assert policy.major_company(name, ['Microsoft','Apple','Alphabet']) == expected


def test_sponsorship_rejection(job):
    assert policy.coarse_reject(dict(job, description='We cannot sponsor visas.'), {'requires_sponsorship': True})
    assert policy.coarse_reject(dict(job, description='No visa sponsorship is available.'), {'requires_sponsorship': True})
    assert not policy.coarse_reject(dict(job, description='No visa sponsorship is available.'), {'requires_sponsorship': False})
    assert not policy.coarse_reject(dict(job, description='Visa sponsorship is available.'))


def test_factual_and_decision_answers(monkeypatch, job):
    profile={'name':'Test Candidate','email':'test@example.org', 'requires_sponsorship': True}
    assert ai.answer('Email', [], profile, job, {}) == 'test@example.org'
    assert ai.answer('Will you require sponsorship in the future?', ['Yes','No'], profile, job, {}) == 'Yes'
    with pytest.raises(policy.NeedsReview):
        ai.answer('What salary do you expect?', [], profile, job, {})
    assert ai.answer('What salary do you expect?', [], profile, job, {'What salary do you expect?':'120000'}) == '120000'
    monkeypatch.setattr(ai, 'structured', lambda *a: ai.Answer(answer='Ten years', needs_decision=False, evidence=['Invented'], reason=''))
    with pytest.raises(policy.NeedsReview):
        ai.answer('Describe your experience', [], profile, job, {})


def test_options_require_an_exact_match(job):
    with pytest.raises(policy.NeedsReview) as exc:
        ai.answer('Will you require sponsorship?', ['I do', 'I do not'], {}, job, {})
    assert exc.value.questions == ['Will you require sponsorship?']


def message(body='Your Example Research verification code is: 123456', sender='noreply@ashbyhq.com', recipient='test@example.org', time=2000):
    import base64
    return {'internalDate':str(time*1000),'payload':{'mimeType':'text/plain','headers':[
        {'name':'From','value':sender},{'name':'To','value':recipient},
        {'name':'Authentication-Results','value':'mx.google.com; dkim=pass header.d=ashbyhq.com'}],
        'body':{'data':base64.urlsafe_b64encode(body.encode()).decode()}}}


def test_code_correlation():
    options=dict(recipient='test@example.org',sender_domains=['ashbyhq.com'],company='Example Research',since=1900)
    assert gmail.extract_code(message(), **options) == '123456'
    assert gmail.extract_code(message(sender='noreply@evilashbyhq.com'), **options) is None
    assert gmail.extract_code(message(recipient='other@example.org'), **options) is None
    assert gmail.extract_code(message(time=1800), **options) is None
    assert gmail.extract_code(message(body='Other company verification code: 123456'), **options) is None
    assert gmail.extract_code(message(body='Example Research verification code: 123456; security code: 654321'), **options) is None


def test_excel_separates_outcomes_and_neutralizes_formulas(job):
    jid = db.add(dict(job, company='=HYPERLINK("https://evil.example")',url='https://jobs.example.org/456'))
    db.update(jid,status='needs_review')
    db.update(job['id'], status='applied', applied_at=db.now())
    path=reports.export()
    workbook=load_workbook(path)
    assert workbook['Applied'].max_row == 2
    assert workbook['Needs Review']['A2'].data_type == 's'
    assert workbook['Needs Review']['A2'].value.startswith('=HYPERLINK')
    assert workbook['Pipeline'].max_row == 1


def test_review_requires_reconciliation_and_csrf(job):
    client=TestClient(server.app)
    db.update(job['id'],status='needs_review',attempted=1)
    url=f'/api/jobs/{job["id"]}/resolve'
    assert client.post(url,json={'action':'retry'}).status_code == 403
    headers={'X-Local-Request':'job-agent'}
    assert client.post(url,json={'action':'retry'},headers=headers).status_code == 409
    assert client.post(url,json={'action':'retry','confirm_not_submitted':True},headers=headers).status_code == 200
    assert db.get(job['id'])['approved'] == 1
    assert client.post(url,json={'action':'retry'},headers=headers).status_code == 409


def test_reject_cross_origin_mutations():
    client=TestClient(server.app)
    response=client.post('/api/commands/discover',json={},headers={'X-Local-Request':'job-agent','Origin':'https://evil.example'})
    assert response.status_code == 403


def test_public_url_blocks_local_targets():
    for url in ['http://example.com', 'https://127.0.0.1', 'https://user:pass@example.com', 'https://localhost']:
        with pytest.raises(ValueError):
            policy.public_url(url)


def test_secret_not_in_state(monkeypatch):
    monkeypatch.setattr(server,'secret',lambda name:'secret-value')
    response=TestClient(server.app).get('/api/state')
    assert response.status_code == 200
    assert 'secret-value' not in response.text


def test_report_ambiguous_send_is_not_repeated(job, monkeypatch):
    class Request:
        def execute(self):
            return {}
    class Messages:
        def list(self,**kw):
            return Request()
    class Api:
        def users(self): return self
        def messages(self): return Messages()
    monkeypatch.setattr(gmail,'service',lambda:Api())
    with db.connection() as c:
        c.execute('INSERT INTO reports VALUES(?,?,?,?)',('2026-09-30','sending',None,db.now()))
    with pytest.raises(policy.NeedsReview):
        gmail.send_report('2026-09-30','unused.xlsx')
