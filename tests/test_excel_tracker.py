from pathlib import Path

from openpyxl import load_workbook

from app.jobs.models import (
    ApplicationMethod,
    ApplicationStatus,
    CompanyRule,
    Job,
)
from app.tracking.database import JobDatabase
from app.tracking.excel_tracker import ExcelTracker


def test_excel_tracker_creation(
    tmp_path: Path,
) -> None:
    database = JobDatabase(
        tmp_path / "test_jobs.db"
    )

    export_path = (
        tmp_path / "Job_Application_Tracker.xlsx"
    )

    tracker = ExcelTracker(
        database,
        export_path,
    )

    result = tracker.generate()

    assert result.exists()


def test_required_sheets(
    tmp_path: Path,
) -> None:
    database = JobDatabase(
        tmp_path / "test_jobs.db"
    )

    export_path = (
        tmp_path / "Job_Application_Tracker.xlsx"
    )

    ExcelTracker(
        database,
        export_path,
    ).generate()

    workbook = load_workbook(export_path)

    expected = {
        "Dashboard",
        "All Jobs",
        "Applied",
        "Manual Queue",
        "Interviews",
        "Rejected",
    }

    assert expected.issubset(
        set(workbook.sheetnames)
    )


def test_job_appears_in_all_jobs(
    tmp_path: Path,
) -> None:
    database = JobDatabase(
        tmp_path / "test_jobs.db"
    )

    job = Job(
        company="Example Company",
        title="Software Engineer",
        url="https://example.com/jobs/1",
        application_url=(
            "https://job-boards.greenhouse.io/example/jobs/1"
        ),
        source="Test",
    )

    database.add_job(job)

    export_path = (
        tmp_path / "Job_Application_Tracker.xlsx"
    )

    ExcelTracker(
        database,
        export_path,
    ).generate()

    workbook = load_workbook(export_path)

    sheet = workbook["All Jobs"]

    assert sheet["A2"].value == "Example Company"
    assert sheet["B2"].value == "Software Engineer"
    assert sheet["L2"].value == "https://example.com/jobs/1"
    assert sheet["M2"].value == (
        "https://job-boards.greenhouse.io/example/jobs/1"
    )


def test_manual_job_appears_in_manual_queue(
    tmp_path: Path,
) -> None:
    database = JobDatabase(
        tmp_path / "test_jobs.db"
    )

    job = Job(
        company="Microsoft",
        title="Software Engineer",
        url="https://example.com/jobs/microsoft",
        source="Test",
        priority_company=True,
        company_rule=CompanyRule.MANUAL,
        application_method=ApplicationMethod.MANUAL,
    )

    database.add_job(job)

    export_path = (
        tmp_path / "Job_Application_Tracker.xlsx"
    )

    ExcelTracker(
        database,
        export_path,
    ).generate()

    workbook = load_workbook(export_path)

    sheet = workbook["Manual Queue"]

    assert sheet["A2"].value == "Microsoft"


def test_applied_job_appears_in_applied_sheet(
    tmp_path: Path,
) -> None:
    database = JobDatabase(
        tmp_path / "test_jobs.db"
    )

    job = Job(
        company="Example Company",
        title="ML Engineer",
        url="https://example.com/jobs/ml",
        source="Test",
        status=ApplicationStatus.APPLIED,
    )

    database.add_job(job)

    export_path = (
        tmp_path / "Job_Application_Tracker.xlsx"
    )

    ExcelTracker(
        database,
        export_path,
    ).generate()

    workbook = load_workbook(export_path)

    sheet = workbook["Applied"]

    assert sheet["A2"].value == "Example Company"


def test_review_job_appears_in_manual_queue(
    tmp_path: Path,
) -> None:
    database = JobDatabase(
        tmp_path / "test_jobs.db"
    )

    job = Job(
        company="Example Startup",
        title="Software Engineer I",
        url="https://example.com/jobs/review",
        source="Test",
        company_rule=CompanyRule.AUTO,
        application_method=ApplicationMethod.REVIEW,
        status=ApplicationStatus.NEEDS_REVIEW,
        fit_score=56.25,
    )

    database.add_job(job)

    export_path = (
        tmp_path / "Job_Application_Tracker.xlsx"
    )

    ExcelTracker(
        database,
        export_path,
    ).generate()

    workbook = load_workbook(export_path)

    sheet = workbook["Manual Queue"]

    assert sheet["A2"].value == "Example Startup"
    assert sheet["F2"].value == 56.25
    assert sheet["G2"].value == CompanyRule.AUTO.value
    assert (
        sheet["H2"].value
        == ApplicationMethod.REVIEW.value
    )
    assert (
        sheet["I2"].value
        == ApplicationStatus.NEEDS_REVIEW.value
    )


def test_dashboard_manual_queue_counts_review_jobs(
    tmp_path: Path,
) -> None:
    database = JobDatabase(
        tmp_path / "test_jobs.db"
    )

    manual_job = Job(
        company="Microsoft",
        title="Software Engineer I",
        url="https://example.com/jobs/manual",
        source="Test",
        company_rule=CompanyRule.MANUAL,
        application_method=ApplicationMethod.MANUAL,
        status=ApplicationStatus.NEEDS_REVIEW,
    )

    review_job = Job(
        company="Example Startup",
        title="Software Engineer I",
        url="https://example.com/jobs/review",
        source="Test",
        company_rule=CompanyRule.AUTO,
        application_method=ApplicationMethod.REVIEW,
        status=ApplicationStatus.NEEDS_REVIEW,
        fit_score=56.25,
    )

    database.add_job(manual_job)
    database.add_job(review_job)

    export_path = (
        tmp_path / "Job_Application_Tracker.xlsx"
    )

    ExcelTracker(
        database,
        export_path,
    ).generate()

    workbook = load_workbook(export_path)

    sheet = workbook["Dashboard"]

    metrics = {
        sheet.cell(row=row, column=1).value:
        sheet.cell(row=row, column=2).value
        for row in range(3, sheet.max_row + 1)
    }

    assert metrics["Manual Queue"] == 2
    assert metrics["Needs Review"] == 2
