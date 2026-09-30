from __future__ import annotations

import json
from pathlib import Path
from threading import Thread
from unittest.mock import MagicMock
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import pytest

from app.dashboard import server as dashboard_server
from app.applications.adapters.detector import ATSProvider
from app.dashboard.application_preview import (
    ApplicationPreviewResult,
    ApplicationPreviewStatus,
)
from app.dashboard.application_launch import (
    ApplicationLaunchResult,
    ApplicationLaunchStatus,
    PreparedApplicationPreview,
)
from app.dashboard.review_sessions import ReviewSessionSnapshot
from app.dashboard.submission_recording import (
    DashboardSubmissionResult,
    DashboardSubmissionStatus,
    PreparedSubmissionRecording,
    SubmissionReviewStatus,
)
from app.dashboard.server import create_dashboard_server
from app.dashboard.review_queue import build_review_queue_item
from app.dashboard.review_resolution import (
    ApplicationReviewResolutionService,
    ReviewResolutionOutcome,
    ReviewResolutionRecord,
    ReviewResolutionResult,
    ReviewResolutionStatus,
)
from app.applications.target_review import (
    ApplicationTargetReviewService,
    TargetReviewResult,
    TargetReviewStatus,
)
from app.jobs.pipeline import PipelineOutcome
from app.dashboard.tracker_reader import (
    DashboardDataError,
    TrackerWorkbookReader,
)
from app.jobs.models import (
    ApplicationMethod,
    ApplicationStatus,
    CompanyRule,
    Job,
    JobCategory,
)
from app.tracking.database import JobDatabase
from app.tracking.excel_tracker import ExcelTracker


def make_workbook(tmp_path: Path) -> Path:
    database = JobDatabase(tmp_path / "jobs.db")
    database.add_job(
        Job(
            company="Example Labs",
            title="Software Engineer I",
            location="Seattle, WA",
            url="https://example.com/jobs/1",
            application_url=(
                "https://job-boards.greenhouse.io/example/jobs/1"
            ),
            source="linkedin_composio",
            category=JobCategory.SDE,
            fit_score=84.5,
            status=ApplicationStatus.NEEDS_APPLICATION,
            application_method=ApplicationMethod.AUTO,
            notes="Ready for the controlled workflow.",
        )
    )
    database.add_job(
        Job(
            company="Microsoft",
            title="Machine Learning Engineer",
            location="Redmond, WA",
            url="https://example.com/jobs/2",
            source="greenhouse",
            category=JobCategory.AI_ML,
            company_rule=CompanyRule.MANUAL,
            application_method=ApplicationMethod.MANUAL,
            status=ApplicationStatus.NEEDS_REVIEW,
            notes="Manual tailoring required.",
        )
    )
    database.add_job(
        Job(
            company="Applied Company",
            title="Data Scientist I",
            location="Remote, US",
            url="https://example.com/jobs/3",
            source="greenhouse",
            category=JobCategory.DATA_SCIENCE,
            application_method=ApplicationMethod.MANUAL,
            status=ApplicationStatus.APPLIED,
        )
    )
    path = tmp_path / "Job_Application_Tracker.xlsx"
    ExcelTracker(database, path).generate()
    return path


def test_reader_returns_dashboard_snapshot(tmp_path: Path) -> None:
    snapshot = TrackerWorkbookReader(make_workbook(tmp_path)).snapshot()

    assert snapshot["source"]["stale"] is False
    assert snapshot["source"]["warning"] is None
    assert snapshot["metrics"] == {
        "total_jobs": 3,
        "applied": 1,
        "manual_queue": 2,
        "review_queue": 1,
        "needs_review": 1,
        "in_progress": 0,
        "offers": 0,
        "rejected": 0,
    }
    assert snapshot["facets"]["statuses"] == [
        "APPLIED",
        "NEEDS_APPLICATION",
        "NEEDS_REVIEW",
    ]
    example_job = next(
        job
        for job in snapshot["jobs"]
        if job["company"] == "Example Labs"
    )
    assert example_job["fit_score"] == 84.5
    assert example_job["job_url"] == (
        "https://example.com/jobs/1"
    )
    assert example_job["application_url"] == (
        "https://job-boards.greenhouse.io/example/jobs/1"
    )
    assert example_job["review_required"] is False
    assert example_job["review_kind"] is None
    assert example_job["review_reason"] is None
    json.dumps(snapshot)


@pytest.mark.parametrize(
    ("job", "kind", "reason_fragment"),
    [
        (
            {
                "status": "NEEDS_REVIEW",
                "application_url": None,
                "application_method": "REVIEW",
            },
            "TARGET_REQUIRED",
            "verified application target",
        ),
        (
            {
                "status": "NEEDS_REVIEW",
                "application_url": "https://example.com/jobs/1",
                "application_method": "REVIEW",
            },
            "INVALID_TARGET",
            "supported ATS",
        ),
        (
            {
                "status": "NEEDS_REVIEW",
                "application_url": (
                    "https://example.wd1.myworkdayjobs.com/en-US/External/"
                    "job/Role_R123/apply/applyManually"
                ),
                "application_method": "REVIEW",
            },
            "WORKDAY_MULTI_STEP",
            "later application steps",
        ),
        (
            {"status": "FORM_STARTED"},
            "FORM_REVIEW",
            "submit manually",
        ),
        (
            {"status": "SUBMISSION_UNCONFIRMED"},
            "SUBMISSION_UNCONFIRMED",
            "confirmation evidence",
        ),
        (
            {"status": "ERROR"},
            "APPLICATION_ERROR",
            "recorded error",
        ),
        (
            {
                "status": "NEEDS_APPLICATION",
                "company_rule": "MANUAL",
                "application_method": "MANUAL",
            },
            "MANUAL_APPLICATION",
            "handled manually",
        ),
    ],
)
def test_review_queue_classifies_actionable_states(
    job: dict[str, object],
    kind: str,
    reason_fragment: str,
) -> None:
    item = build_review_queue_item(job)

    assert item["review_required"] is True
    assert item["review_kind"] == kind
    assert reason_fragment in item["review_reason"]


@pytest.mark.parametrize(
    "status",
    ["APPLIED", "REJECTED", "OFFER", "WITHDRAWN", "FILTERED_OUT"],
)
def test_review_queue_excludes_completed_or_inactive_states(status: str) -> None:
    assert build_review_queue_item({"status": status}) == {
        "review_required": False,
        "review_kind": None,
        "review_reason": None,
    }


def test_reader_reuses_snapshot_until_file_changes(tmp_path: Path) -> None:
    reader = TrackerWorkbookReader(make_workbook(tmp_path))
    first = reader.snapshot()
    second = reader.snapshot()

    assert first is second


def test_reader_serves_last_good_snapshot_when_reload_fails(
    tmp_path: Path,
) -> None:
    path = make_workbook(tmp_path)
    reader = TrackerWorkbookReader(path)
    first = reader.snapshot()

    path.write_bytes(b"not an xlsx workbook")
    stale = reader.snapshot()

    assert stale is not first
    assert stale["jobs"] == first["jobs"]
    assert stale["source"]["stale"] is True
    assert "last successful" in stale["source"]["warning"].lower()


def test_missing_workbook_is_actionable(tmp_path: Path) -> None:
    reader = TrackerWorkbookReader(tmp_path / "missing.xlsx")

    with pytest.raises(DashboardDataError, match="does not exist"):
        reader.snapshot()


def test_reader_rejects_missing_all_jobs_sheet(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "invalid.xlsx"
    Workbook().save(path)

    with pytest.raises(DashboardDataError, match="All Jobs"):
        TrackerWorkbookReader(path).snapshot()


def test_reader_rejects_unexpected_headers(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "invalid.xlsx"
    workbook = Workbook()
    workbook.active.title = "All Jobs"
    workbook.active.append(["Company", "Unexpected"])
    workbook.save(path)

    with pytest.raises(DashboardDataError, match="headers"):
        TrackerWorkbookReader(path).snapshot()


def test_non_web_job_url_is_not_exposed_as_link(tmp_path: Path) -> None:
    path = make_workbook(tmp_path)
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    workbook["All Jobs"]["L2"] = "javascript:alert(1)"
    workbook.save(path)

    snapshot = TrackerWorkbookReader(path).snapshot()
    applied_job = next(
        job
        for job in snapshot["jobs"]
        if job["company"] == "Applied Company"
    )
    assert applied_job["job_url"] is None


def test_http_server_exposes_ui_api_and_health(tmp_path: Path) -> None:
    reader = TrackerWorkbookReader(make_workbook(tmp_path))
    server = create_dashboard_server(
        reader=reader,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base_url = f"http://127.0.0.1:{server.server_port}"

    try:
        with urlopen(base_url + "/", timeout=5) as response:
            html = response.read().decode("utf-8")
            assert response.status == 200
            assert "Job Application Dashboard" in html
            assert 'id="auto-prepare-consent"' in html
            assert "open and fill every eligible application" in html

        with urlopen(base_url + "/api/jobs", timeout=5) as response:
            payload = json.load(response)
            assert response.status == 200
            assert payload["metrics"]["total_jobs"] == 3
            assert response.headers["Cache-Control"] == "no-store"

        with urlopen(base_url + "/api/health", timeout=5) as response:
            payload = json.load(response)
            assert payload["status"] == "ok"
            assert payload["workbook_available"] is True

        with urlopen(base_url + "/static/app.js", timeout=5) as response:
            script = response.read().decode("utf-8")
            assert "setInterval" in script
            assert "application-target" in script
            assert "Set target" in script
            assert "review_reason" in script
            assert "review-filter" in script
            assert "review-resolution" in script
            assert "review-history" in script
            assert "history-outcome-filter" in script
            assert "history-kind-filter" in script
            assert "application-preview" in script
            assert "Preview" in script
            assert "application-launch" in script
            assert "AUTHORIZE_EXTERNAL_BROWSER" in script
            assert "Authorize and open application" in script
            assert "review-session/close" in script
            assert "Active review session" in script
            assert "Close review session" in script
            assert "submission-review" in script
            assert "submission-recording" in script
            assert "Record result" in script
            assert "independently observed" in script
            assert "autoPrepareConsent" in script
            assert "await prepareEligibleApplications()" in script
            assert "innerHTML" not in script
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_http_server_reports_missing_workbook(tmp_path: Path) -> None:
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(tmp_path / "missing.xlsx"),
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        request = Request(
            f"http://127.0.0.1:{server.server_port}/api/jobs"
        )
        with pytest.raises(HTTPError) as error:
            urlopen(request, timeout=5)
        assert error.value.code == 503
        payload = json.loads(error.value.read().decode("utf-8"))
        assert "does not exist" in payload["error"]
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_http_server_rejects_writes(tmp_path: Path) -> None:
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        request = Request(
            f"http://127.0.0.1:{server.server_port}/api/jobs",
            method="POST",
            data=b"{}",
        )
        with pytest.raises(HTTPError) as error:
            urlopen(request, timeout=5)
        assert error.value.code == 405
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_target_review_api_assigns_explicit_url(tmp_path: Path) -> None:
    reviewer = MagicMock()
    reviewer.assign.return_value = TargetReviewResult(
        status=TargetReviewStatus.UPDATED,
        reason="Reprocessed.",
        job_id=7,
        application_url=(
            "https://job-boards.greenhouse.io/example/jobs/123"
        ),
        pipeline_outcome=PipelineOutcome.AUTO_READY,
        export_path=tmp_path / "tracker.xlsx",
    )
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        target_review_service=reviewer,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        body = json.dumps({
            "application_url": (
                "https://job-boards.greenhouse.io/example/jobs/123"
            )
        }).encode("utf-8")
        request = Request(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/application-target",
            method="POST",
            data=body,
            headers={"Content-Type": "application/json"},
        )
        with urlopen(request, timeout=5) as response:
            payload = json.load(response)
        assert response.status == 200
        assert payload["status"] == "UPDATED"
        assert payload["pipeline_outcome"] == "AUTO_READY"
        reviewer.assign.assert_called_once_with(
            job_id=7,
            application_url=(
                "https://job-boards.greenhouse.io/example/jobs/123"
            ),
        )
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_editable_jobs_snapshot_exposes_only_eligible_record_id(
    tmp_path: Path,
) -> None:
    database = JobDatabase(tmp_path / "review.db")
    job_id = database.add_job(Job(
        company="Review Company",
        title="Software Engineer I",
        location="Seattle, WA",
        url="https://www.linkedin.com/jobs/view/98765",
        source="linkedin_composio",
        status=ApplicationStatus.NEEDS_REVIEW,
        application_method=ApplicationMethod.REVIEW,
    ))
    workbook = tmp_path / "review.xlsx"
    tracker = ExcelTracker(database, workbook)
    tracker.generate()
    reviewer = ApplicationTargetReviewService(
        database=database,
        pipeline=MagicMock(),
        tracker=tracker,
    )
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(workbook),
        target_review_service=reviewer,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        with urlopen(
            f"http://127.0.0.1:{server.server_port}/api/jobs",
            timeout=5,
        ) as response:
            payload = json.load(response)
        assert payload["jobs"][0]["job_id"] == job_id
        assert payload["jobs"][0]["target_review_eligible"] is True
        assert payload["jobs"][0]["review_required"] is True
        assert payload["jobs"][0]["review_kind"] == "TARGET_REQUIRED"
        assert "verified application target" in (
            payload["jobs"][0]["review_reason"]
        )
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_target_review_api_rejects_non_json_body(tmp_path: Path) -> None:
    reviewer = MagicMock()
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        target_review_service=reviewer,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        request = Request(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/application-target",
            method="POST",
            data=b"application_url=https://example.com",
            headers={"Content-Type": "text/plain"},
        )
        with pytest.raises(HTTPError) as error:
            urlopen(request, timeout=5)
        assert error.value.code == 415
        reviewer.assign.assert_not_called()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_review_resolution_api_records_explicit_decision(
    tmp_path: Path,
) -> None:
    service = MagicMock()
    service.record.return_value = ReviewResolutionResult(
        status=ReviewResolutionStatus.RECORDED,
        reason="Review decision recorded.",
        job_id=7,
        record=MagicMock(
            outcome=ReviewResolutionOutcome.DEFERRED,
            review_kind="TARGET_REQUIRED",
            note="Waiting for the employer application link.",
            created_at="2026-09-29T12:00:00",
        ),
        export_path=tmp_path / "tracker.xlsx",
    )
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        review_resolution_service=service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        request = Request(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/review-resolution",
            method="POST",
            data=json.dumps({
                "outcome": "DEFERRED",
                "note": "Waiting for the employer application link.",
            }).encode("utf-8"),
            headers={"Content-Type": "application/json"},
        )
        with urlopen(request, timeout=5) as response:
            payload = json.load(response)

        assert response.status == 200
        assert payload["status"] == "RECORDED"
        assert payload["outcome"] == "DEFERRED"
        service.record.assert_called_once_with(
            job_id=7,
            outcome="DEFERRED",
            note="Waiting for the employer application link.",
        )
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_dashboard_applies_latest_resolution_without_changing_job_status(
    tmp_path: Path,
) -> None:
    database = JobDatabase(tmp_path / "review-resolution.db")
    job_id = database.add_job(Job(
        company="Review Company",
        title="Software Engineer I",
        location="Seattle, WA",
        url="https://www.linkedin.com/jobs/view/555",
        source="linkedin_composio",
        status=ApplicationStatus.NEEDS_REVIEW,
        application_method=ApplicationMethod.REVIEW,
    ))
    workbook = tmp_path / "review-resolution.xlsx"
    tracker = ExcelTracker(database, workbook)
    tracker.generate()
    service = ApplicationReviewResolutionService(
        database=database,
        tracker=tracker,
    )
    service.record(
        job_id=job_id,
        outcome="DISMISSED",
        note="The role closed before an application target was published.",
    )
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(workbook),
        review_resolution_service=service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        with urlopen(
            f"http://127.0.0.1:{server.server_port}/api/jobs",
            timeout=5,
        ) as response:
            payload = json.load(response)

        item = payload["jobs"][0]
        assert payload["metrics"]["review_queue"] == 0
        assert item["review_required"] is False
        assert item["review_outcome"] == "DISMISSED"
        assert "role closed" in item["review_note"]
        assert database.get_job_by_id(job_id).status == (
            ApplicationStatus.NEEDS_REVIEW
        )
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_review_history_api_returns_chronological_read_only_records(
    tmp_path: Path,
) -> None:
    service = MagicMock()
    service.history.return_value = [
        ReviewResolutionRecord(
            resolution_id=1,
            job_id=7,
            review_kind="TARGET_REQUIRED",
            outcome=ReviewResolutionOutcome.DEFERRED,
            note="Waiting for an application link.",
            created_at="2026-09-29T12:00:00",
        ),
        ReviewResolutionRecord(
            resolution_id=2,
            job_id=7,
            review_kind="TARGET_REQUIRED",
            outcome=ReviewResolutionOutcome.RESOLVED,
            note="Applied manually.",
            created_at="2026-09-29T13:00:00",
        ),
    ]
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        review_resolution_service=service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        with urlopen(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/review-history",
            timeout=5,
        ) as response:
            payload = json.load(response)

        assert response.status == 200
        assert payload["job_id"] == 7
        assert [item["outcome"] for item in payload["history"]] == [
            "DEFERRED",
            "RESOLVED",
        ]
        assert payload["history"][1]["note"] == "Applied manually."
        service.history.assert_called_once_with(7)
        service.record.assert_not_called()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_application_preview_api_is_read_only_and_never_runs_a_browser(
    tmp_path: Path,
) -> None:
    service = MagicMock()
    service.preview.return_value = ApplicationPreviewResult(
        status=ApplicationPreviewStatus.READY,
        reason="Job passed browser-free application preflight.",
        job_id=7,
        job=Job(
            company="Example",
            title="Software Engineer I",
            location="Seattle, WA",
            url="https://www.linkedin.com/jobs/view/123",
            application_url=(
                "https://job-boards.greenhouse.io/example/jobs/123"
            ),
            source="test",
            status=ApplicationStatus.NEEDS_APPLICATION,
            company_rule=CompanyRule.AUTO,
            application_method=ApplicationMethod.AUTO,
            resume_used="data/resumes/sde_resume.pdf",
            fit_score=92,
        ),
        source_url="https://www.linkedin.com/jobs/view/123",
        application_url=(
            "https://job-boards.greenhouse.io/example/jobs/123"
        ),
        provider=ATSProvider.GREENHOUSE,
        resume="data/resumes/sde_resume.pdf",
    )
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        application_preview_service=service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        with urlopen(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/application-preview",
            timeout=5,
        ) as response:
            payload = json.load(response)

        assert response.status == 200
        assert payload["status"] == "READY"
        assert payload["job"]["company"] == "Example"
        assert payload["source_url"] == (
            "https://www.linkedin.com/jobs/view/123"
        )
        assert payload["application_url"] == (
            "https://job-boards.greenhouse.io/example/jobs/123"
        )
        assert payload["ats_provider"] == "GREENHOUSE"
        assert payload["resume"] == "data/resumes/sde_resume.pdf"
        assert payload["safety"] == {
            "browser_started": False,
            "workflow_ran": False,
            "external_authorization_required": True,
            "fields_filled": False,
            "files_uploaded": False,
            "may_submit": False,
        }
        service.preview.assert_called_once_with(7)
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_preview_issues_authorization_without_launching_workflow(
    tmp_path: Path,
) -> None:
    preview = ApplicationPreviewResult(
        status=ApplicationPreviewStatus.READY,
        reason="Browser-free preflight passed.",
        job_id=7,
    )
    launch_service = MagicMock()
    launch_service.prepare.return_value = PreparedApplicationPreview(
        preview=preview,
        authorization_token="exact-preview-token",
        authorization_expires_in_seconds=120,
        review_session=ReviewSessionSnapshot(
            job_id=7,
            target_url="https://job-boards.greenhouse.io/example/jobs/123",
            active=True,
            expires_in_seconds=900,
        ),
    )
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        application_launch_service=launch_service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        with urlopen(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/application-preview",
            timeout=5,
        ) as response:
            payload = json.load(response)

        assert payload["authorization"] == {
            "token": "exact-preview-token",
            "expires_in_seconds": 120,
            "confirmation": "AUTHORIZE_EXTERNAL_BROWSER",
        }
        assert payload["review_session"]["active"] is True
        assert payload["review_session"]["expires_in_seconds"] == 900
        launch_service.prepare.assert_called_once_with(7)
        launch_service.launch.assert_not_called()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_application_launch_api_requires_exact_explicit_authorization(
    tmp_path: Path,
) -> None:
    launch_service = MagicMock()
    launch_service.launch.return_value = ApplicationLaunchResult(
        status=ApplicationLaunchStatus.READY_FOR_REVIEW,
        reason="Authorized fields populated; human review required.",
        job_id=7,
        completed_actions=4,
        application_status=ApplicationStatus.FORM_STARTED,
        export_path=tmp_path / "tracker.xlsx",
        review_session=ReviewSessionSnapshot(
            job_id=7,
            target_url="https://job-boards.greenhouse.io/example/jobs/123",
            active=True,
            expires_in_seconds=900,
        ),
    )
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        application_launch_service=launch_service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        body = json.dumps({
            "authorization_token": "exact-preview-token",
            "confirmation": "AUTHORIZE_EXTERNAL_BROWSER",
        }).encode("utf-8")
        request = Request(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/application-launch",
            method="POST",
            data=body,
            headers={"Content-Type": "application/json"},
        )
        with urlopen(request, timeout=5) as response:
            payload = json.load(response)

        assert response.status == 200
        assert payload["status"] == "READY_FOR_REVIEW"
        assert payload["completed_actions"] == 4
        assert payload["application_status"] == "FORM_STARTED"
        assert payload["may_submit"] is False
        assert payload["review_session"]["active"] is True
        launch_service.launch.assert_called_once_with(
            job_id=7,
            authorization_token="exact-preview-token",
            confirmation="AUTHORIZE_EXTERNAL_BROWSER",
        )
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_review_session_close_api_requires_explicit_confirmation(
    tmp_path: Path,
) -> None:
    launch_service = MagicMock()
    launch_service.close_review_session.return_value = True
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        application_launch_service=launch_service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        request = Request(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/review-session/close",
            method="POST",
            data=json.dumps({
                "confirmation": "CLOSE_REVIEW_SESSION",
            }).encode("utf-8"),
            headers={"Content-Type": "application/json"},
        )
        with urlopen(request, timeout=5) as response:
            payload = json.load(response)

        assert response.status == 200
        assert payload == {
            "status": "CLOSED",
            "job_id": 7,
            "active": False,
        }
        launch_service.close_review_session.assert_called_once_with(7)
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_review_session_close_rejects_missing_confirmation(
    tmp_path: Path,
) -> None:
    launch_service = MagicMock()
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        application_launch_service=launch_service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        request = Request(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/review-session/close",
            method="POST",
            data=b"{}",
            headers={"Content-Type": "application/json"},
        )
        with pytest.raises(HTTPError) as error:
            urlopen(request, timeout=5)

        assert error.value.code == 400
        launch_service.close_review_session.assert_not_called()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_dashboard_server_close_cleans_all_review_sessions(
    tmp_path: Path,
) -> None:
    launch_service = MagicMock()
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        application_launch_service=launch_service,
        host="127.0.0.1",
        port=0,
    )

    server.server_close()

    launch_service.close_all_review_sessions.assert_called_once_with()


def test_application_launch_api_rejects_missing_authorization(
    tmp_path: Path,
) -> None:
    launch_service = MagicMock()
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        application_launch_service=launch_service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        request = Request(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/application-launch",
            method="POST",
            data=b"{}",
            headers={"Content-Type": "application/json"},
        )
        with pytest.raises(HTTPError) as error:
            urlopen(request, timeout=5)

        assert error.value.code == 400
        launch_service.launch.assert_not_called()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_submission_review_api_is_read_only_and_issues_exact_authorization(
    tmp_path: Path,
) -> None:
    service = MagicMock()
    service.prepare.return_value = PreparedSubmissionRecording(
        status=SubmissionReviewStatus.READY,
        reason="Record the independently observed result.",
        job_id=7,
        job=Job(
            company="Example",
            title="Software Engineer I",
            location="Seattle, WA",
            url="https://example.com/jobs/7",
            source="test",
            status=ApplicationStatus.FORM_STARTED,
        ),
        authorization_token="record-token",
        authorization_expires_in_seconds=300,
    )
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        submission_recording_service=service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        with urlopen(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/submission-review",
            timeout=5,
        ) as response:
            payload = json.load(response)

        assert response.status == 200
        assert payload["status"] == "READY"
        assert payload["job"]["application_status"] == "FORM_STARTED"
        assert payload["authorization"] == {
            "token": "record-token",
            "expires_in_seconds": 300,
            "confirmations": {
                "CONFIRMED": "RECORD_CONFIRMED_SUBMISSION",
                "UNCONFIRMED": "RECORD_UNCONFIRMED_SUBMISSION",
                "NOT_SUBMITTED": "RECORD_NOT_SUBMITTED",
            },
        }
        assert payload["may_submit"] is False
        service.prepare.assert_called_once_with(7)
        service.record.assert_not_called()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_submission_recording_api_forwards_explicit_outcome(
    tmp_path: Path,
) -> None:
    service = MagicMock()
    service.record.return_value = DashboardSubmissionResult(
        status=DashboardSubmissionStatus.CONFIRMED,
        reason="Successful submission was independently confirmed.",
        job_id=7,
        job=Job(
            company="Example",
            title="Software Engineer I",
            location="Seattle, WA",
            url="https://example.com/jobs/7",
            source="test",
            status=ApplicationStatus.APPLIED,
        ),
        export_path=tmp_path / "tracker.xlsx",
    )
    server = create_dashboard_server(
        reader=TrackerWorkbookReader(make_workbook(tmp_path)),
        submission_recording_service=service,
        host="127.0.0.1",
        port=0,
    )
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        request = Request(
            f"http://127.0.0.1:{server.server_port}"
            "/api/jobs/7/submission-recording",
            method="POST",
            data=json.dumps({
                "authorization_token": "record-token",
                "outcome": "CONFIRMED",
                "confirmation": "RECORD_CONFIRMED_SUBMISSION",
                "evidence": "Portal confirmation number 123",
            }).encode("utf-8"),
            headers={"Content-Type": "application/json"},
        )
        with urlopen(request, timeout=5) as response:
            payload = json.load(response)

        assert response.status == 200
        assert payload["status"] == "CONFIRMED"
        assert payload["job"]["application_status"] == "APPLIED"
        assert payload["may_submit"] is False
        service.record.assert_called_once_with(
            job_id=7,
            authorization_token="record-token",
            outcome="CONFIRMED",
            confirmation="RECORD_CONFIRMED_SUBMISSION",
            evidence="Portal confirmation number 123",
        )
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_dashboard_main_composes_preview_and_launch_services(
    tmp_path: Path,
    monkeypatch,
) -> None:
    database = MagicMock()
    tracker = MagicMock()
    preview_service = MagicMock()
    launch_service = MagicMock()
    submission_runner = MagicMock()
    submission_service = MagicMock()
    review_session_manager = MagicMock()
    http_server = MagicMock(server_port=8765)
    database_builder = MagicMock(return_value=database)
    tracker_builder = MagicMock(return_value=tracker)
    preview_builder = MagicMock(return_value=preview_service)
    server_builder = MagicMock(return_value=http_server)
    monkeypatch.setattr(dashboard_server, "JobDatabase", database_builder)
    monkeypatch.setattr(dashboard_server, "ExcelTracker", tracker_builder)
    monkeypatch.setattr(
        dashboard_server,
        "ApplicationPreviewService",
        preview_builder,
    )
    launch_builder = MagicMock(return_value=launch_service)
    monkeypatch.setattr(
        dashboard_server,
        "DashboardApplicationLaunchService",
        launch_builder,
    )
    session_manager_builder = MagicMock(
        return_value=review_session_manager
    )
    monkeypatch.setattr(
        dashboard_server,
        "ApplicationReviewSessionManager",
        session_manager_builder,
    )
    launcher_builder = MagicMock(return_value=MagicMock())
    monkeypatch.setattr(
        dashboard_server,
        "build_single_job_application_launcher_from_dependencies",
        launcher_builder,
    )
    submission_runner_builder = MagicMock(
        return_value=submission_runner
    )
    monkeypatch.setattr(
        dashboard_server,
        "build_submission_recording_runner_from_dependencies",
        submission_runner_builder,
    )
    submission_service_builder = MagicMock(
        return_value=submission_service
    )
    monkeypatch.setattr(
        dashboard_server,
        "DashboardSubmissionRecordingService",
        submission_service_builder,
    )
    monkeypatch.setattr(
        dashboard_server,
        "create_dashboard_server",
        server_builder,
    )
    monkeypatch.setattr(
        dashboard_server,
        "build_job_pipeline",
        MagicMock(),
    )

    exit_code = dashboard_server.main([
        "--database",
        str(tmp_path / "jobs.db"),
        "--workbook",
        str(tmp_path / "tracker.xlsx"),
    ])

    assert exit_code == 0
    preview_builder.assert_called_once_with(database=database)
    assert launcher_builder.call_args.kwargs["database"] is database
    assert launcher_builder.call_args.kwargs["tracker"] is tracker
    assert launcher_builder.call_args.kwargs[
        "execution_session_factory"
    ] is review_session_manager.create_execution_session
    session_manager_builder.assert_called_once_with(
        review_ttl_seconds=900,
    )
    launch_builder.assert_called_once_with(
        preview_service=preview_service,
        launcher=launcher_builder.return_value,
        review_session_manager=review_session_manager,
    )
    submission_runner_builder.assert_called_once_with(
        database=database,
        tracker=tracker,
    )
    submission_service_builder.assert_called_once_with(
        runner=submission_runner,
        active_session_lookup=review_session_manager.snapshot,
    )
    assert server_builder.call_args.kwargs[
        "application_preview_service"
    ] is preview_service
    assert server_builder.call_args.kwargs[
        "application_launch_service"
    ] is launch_service
    assert server_builder.call_args.kwargs[
        "submission_recording_service"
    ] is submission_service
    http_server.serve_forever.assert_called_once_with()
    http_server.server_close.assert_called_once_with()


def test_editable_dashboard_requires_loopback_binding(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="loopback"):
        create_dashboard_server(
            reader=TrackerWorkbookReader(make_workbook(tmp_path)),
            target_review_service=MagicMock(),
            host="0.0.0.0",
            port=0,
        )
