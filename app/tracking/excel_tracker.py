from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.jobs.models import (
    ApplicationMethod,
    ApplicationStatus,
    CompanyRule,
    Job,
)
from app.tracking.database import JobDatabase


DEFAULT_EXPORT_PATH = Path(
    "data/exports/Job_Application_Tracker.xlsx"
)


class ExcelTracker:
    """
    Generates a human-readable Excel workbook from SQLite.

    SQLite remains the authoritative source of truth.
    This workbook is a reporting/export layer only.
    """

    JOB_HEADERS = [
        "Company",
        "Title",
        "Location",
        "Category",
        "Seniority",
        "Fit Score",
        "Company Rule",
        "Application Method",
        "Status",
        "Resume Used",
        "Source",
        "Job URL",
        "Application URL",
        "Date Posted",
        "Date Found",
        "Date Applied",
        "Notes",
    ]

    def __init__(
        self,
        database: JobDatabase,
        export_path: Path | str = DEFAULT_EXPORT_PATH,
    ) -> None:
        self.database = database
        self.export_path = Path(export_path)

        self.export_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

    # ---------------------------------------------------------
    # Public API
    # ---------------------------------------------------------

    def generate(self) -> Path:
        """
        Generate the complete Excel tracker from SQLite.
        """

        jobs = self.database.get_all_jobs()

        workbook = Workbook()

        # Remove the default worksheet.
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        self._create_dashboard(workbook, jobs)

        self._create_job_sheet(
            workbook,
            "All Jobs",
            jobs,
        )

        self._create_job_sheet(
            workbook,
            "Applied",
            [
                job
                for job in jobs
                if job.status == ApplicationStatus.APPLIED
            ],
        )

        self._create_job_sheet(
            workbook,
            "Manual Queue",
            [
                job
                for job in jobs
                if self._belongs_in_manual_queue(job)
            ],
        )

        self._create_job_sheet(
            workbook,
            "Interviews",
            [
                job
                for job in jobs
                if job.status
                in {
                    ApplicationStatus.OA,
                    ApplicationStatus.INTERVIEW,
                    ApplicationStatus.OFFER,
                }
            ],
        )

        self._create_job_sheet(
            workbook,
            "Rejected",
            [
                job
                for job in jobs
                if job.status
                == ApplicationStatus.REJECTED
            ],
        )

        workbook.save(self.export_path)

        return self.export_path

    @staticmethod
    def _belongs_in_manual_queue(job: Job) -> bool:
        """
        Return True when a job requires human attention.

        MANUAL represents company-policy routing, such as priority
        companies that must be tailored and applied to manually.

        REVIEW represents fit-gate routing where the job's match
        requires a human decision before continuing.
        """

        return (
            job.company_rule == CompanyRule.MANUAL
            or job.application_method
            in {
                ApplicationMethod.MANUAL,
                ApplicationMethod.REVIEW,
            }
        )

    # ---------------------------------------------------------
    # Dashboard
    # ---------------------------------------------------------

    def _create_dashboard(
        self,
        workbook: Workbook,
        jobs: list[Job],
    ) -> None:
        sheet = workbook.create_sheet("Dashboard")

        sheet["A1"] = "Job Application Dashboard"
        sheet["A1"].font = Font(
            bold=True,
            size=18,
        )

        sheet.merge_cells("A1:B1")

        total_jobs = len(jobs)

        applied = sum(
            job.status == ApplicationStatus.APPLIED
            for job in jobs
        )

        manual_queue = sum(
            self._belongs_in_manual_queue(job)
            for job in jobs
        )

        interviews = sum(
            job.status
            in {
                ApplicationStatus.OA,
                ApplicationStatus.INTERVIEW,
            }
            for job in jobs
        )

        offers = sum(
            job.status == ApplicationStatus.OFFER
            for job in jobs
        )

        rejected = sum(
            job.status == ApplicationStatus.REJECTED
            for job in jobs
        )

        needs_review = sum(
            job.status == ApplicationStatus.NEEDS_REVIEW
            for job in jobs
        )

        metrics = [
            ("Total Jobs", total_jobs),
            ("Applied", applied),
            ("Manual Queue", manual_queue),
            ("Needs Review", needs_review),
            ("OA / Interviews", interviews),
            ("Offers", offers),
            ("Rejected", rejected),
        ]

        start_row = 3

        for row_number, (label, value) in enumerate(
            metrics,
            start=start_row,
        ):
            sheet.cell(
                row=row_number,
                column=1,
                value=label,
            )

            sheet.cell(
                row=row_number,
                column=2,
                value=value,
            )

        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 15

        for row in range(
            start_row,
            start_row + len(metrics),
        ):
            sheet.cell(row=row, column=1).font = Font(
                bold=True
            )

            sheet.cell(
                row=row,
                column=2,
            ).alignment = Alignment(
                horizontal="center"
            )

    # ---------------------------------------------------------
    # Job sheets
    # ---------------------------------------------------------

    def _create_job_sheet(
        self,
        workbook: Workbook,
        sheet_name: str,
        jobs: Iterable[Job],
    ) -> None:
        sheet = workbook.create_sheet(sheet_name)

        jobs = list(jobs)

        for column_number, header in enumerate(
            self.JOB_HEADERS,
            start=1,
        ):
            cell = sheet.cell(
                row=1,
                column=column_number,
                value=header,
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row_number, job in enumerate(
            jobs,
            start=2,
        ):
            values = self._job_to_row(job)

            for column_number, value in enumerate(
                values,
                start=1,
            ):
                sheet.cell(
                    row=row_number,
                    column=column_number,
                    value=value,
                )

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = (
            f"A1:Q{max(sheet.max_row, 1)}"
        )

        self._set_column_widths(sheet)

        # Excel tables require at least one data row.
        if jobs:
            table_name = (
                sheet_name
                .replace(" ", "")
                .replace("-", "")
                + "Table"
            )

            table = Table(
                displayName=table_name,
                ref=f"A1:Q{sheet.max_row}",
            )

            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            table.tableStyleInfo = style

            sheet.add_table(table)

    # ---------------------------------------------------------
    # Conversion
    # ---------------------------------------------------------

    @staticmethod
    def _job_to_row(job: Job) -> list:
        return [
            job.company,
            job.title,
            job.location,
            job.category.value,
            job.seniority.value,
            job.fit_score,
            job.company_rule.value,
            job.application_method.value,
            job.status.value,
            job.resume_used,
            job.source,
            job.url.encoded_string(),
            (
                job.application_url.encoded_string()
                if job.application_url
                else None
            ),
            (
                job.date_posted.isoformat()
                if job.date_posted
                else None
            ),
            job.date_found.isoformat(
                sep=" ",
                timespec="seconds",
            ),
            (
                job.date_applied.isoformat(
                    sep=" ",
                    timespec="seconds",
                )
                if job.date_applied
                else None
            ),
            job.notes,
        ]

    # ---------------------------------------------------------
    # Formatting
    # ---------------------------------------------------------

    @staticmethod
    def _set_column_widths(sheet) -> None:
        widths = {
            1: 24,
            2: 34,
            3: 24,
            4: 18,
            5: 18,
            6: 12,
            7: 16,
            8: 20,
            9: 24,
            10: 32,
            11: 18,
            12: 55,
            13: 55,
            14: 16,
            15: 22,
            16: 22,
            17: 45,
        }

        for column_number, width in widths.items():
            column_letter = get_column_letter(
                column_number
            )

            sheet.column_dimensions[
                column_letter
            ].width = width
