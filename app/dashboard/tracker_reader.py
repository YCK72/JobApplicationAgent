from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from threading import RLock
from typing import Any
from urllib.parse import urlsplit
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.tracking.excel_tracker import ExcelTracker


DEFAULT_WORKBOOK_PATH = Path(
    "data/exports/Job_Application_Tracker.xlsx"
)


class DashboardDataError(RuntimeError):
    """Raised when no safe dashboard snapshot can be loaded."""


class TrackerWorkbookReader:
    """Read and cache the exported workbook without modifying it.

    The tracker may be replaced while the dashboard is polling it. When a
    reload fails after one successful read, the last good snapshot remains
    visible with a stale-data warning until a later reload succeeds.
    """

    HEADER_KEYS = {
        "Company": "company",
        "Title": "title",
        "Location": "location",
        "Category": "category",
        "Seniority": "seniority",
        "Fit Score": "fit_score",
        "Company Rule": "company_rule",
        "Application Method": "application_method",
        "Status": "status",
        "Source": "source",
        "Job URL": "job_url",
        "Application URL": "application_url",
        "Date Posted": "date_posted",
        "Date Found": "date_found",
        "Date Applied": "date_applied",
        "Notes": "notes",
    }

    def __init__(self, workbook_path: Path | str) -> None:
        self.workbook_path = Path(workbook_path)
        self._lock = RLock()
        self._snapshot: dict[str, Any] | None = None
        self._snapshot_signature: tuple[int, int] | None = None
        self._failed_signature: tuple[int, int] | None = None
        self._stale_snapshot: dict[str, Any] | None = None

    def snapshot(self) -> dict[str, Any]:
        with self._lock:
            try:
                signature = self._signature()
            except DashboardDataError as exc:
                return self._stale_or_raise(exc, None)

            if (
                self._snapshot is not None
                and signature == self._snapshot_signature
            ):
                return self._snapshot

            if (
                self._stale_snapshot is not None
                and signature == self._failed_signature
            ):
                return self._stale_snapshot

            try:
                snapshot = self._load(signature)
            except (
                DashboardDataError,
                InvalidFileException,
                BadZipFile,
                OSError,
                ValueError,
            ) as exc:
                error = (
                    exc
                    if isinstance(exc, DashboardDataError)
                    else DashboardDataError(
                        "Tracker workbook is temporarily unavailable."
                    )
                )
                return self._stale_or_raise(error, signature)

            self._snapshot = snapshot
            self._snapshot_signature = signature
            self._failed_signature = None
            self._stale_snapshot = None
            return snapshot

    def _signature(self) -> tuple[int, int]:
        try:
            stat = self.workbook_path.stat()
        except FileNotFoundError:
            raise DashboardDataError(
                f"Tracker workbook does not exist: {self.workbook_path}"
            ) from None
        except OSError:
            raise DashboardDataError(
                "Tracker workbook cannot be accessed."
            ) from None

        if not self.workbook_path.is_file():
            raise DashboardDataError(
                f"Tracker workbook is not a file: {self.workbook_path}"
            )

        return stat.st_mtime_ns, stat.st_size

    def _load(self, signature: tuple[int, int]) -> dict[str, Any]:
        try:
            workbook_bytes = self.workbook_path.read_bytes()
        except OSError:
            raise DashboardDataError(
                "Tracker workbook cannot be read right now."
            ) from None

        workbook = load_workbook(
            BytesIO(workbook_bytes),
            read_only=True,
            data_only=True,
        )
        try:
            if "All Jobs" not in workbook.sheetnames:
                raise DashboardDataError(
                    "Tracker workbook is missing the 'All Jobs' sheet."
                )

            sheet = workbook["All Jobs"]
            rows = sheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                raise DashboardDataError(
                    "The 'All Jobs' sheet is empty."
                ) from None

            headers = list(header_row)
            while headers and headers[-1] is None:
                headers.pop()

            expected_headers = list(ExcelTracker.JOB_HEADERS)
            if headers != expected_headers:
                raise DashboardDataError(
                    "The 'All Jobs' sheet headers do not match the "
                    "tracker contract. Regenerate the workbook."
                )

            jobs = [
                self._job_from_row(header_row, row)
                for row in rows
                if any(value is not None for value in row)
            ]
        finally:
            workbook.close()

        modified_at = datetime.fromtimestamp(
            signature[0] / 1_000_000_000,
            tz=timezone.utc,
        ).isoformat()

        return {
            "source": {
                "name": self.workbook_path.name,
                "modified_at": modified_at,
                "loaded_at": datetime.now(timezone.utc).isoformat(),
                "stale": False,
                "warning": None,
            },
            "metrics": self._metrics(jobs),
            "facets": {
                "statuses": self._facet(jobs, "status"),
                "categories": self._facet(jobs, "category"),
                "methods": self._facet(
                    jobs,
                    "application_method",
                ),
                "sources": self._facet(jobs, "source"),
            },
            "jobs": jobs,
        }

    def _stale_or_raise(
        self,
        error: DashboardDataError,
        signature: tuple[int, int] | None,
    ) -> dict[str, Any]:
        if self._snapshot is None:
            raise error

        stale = deepcopy(self._snapshot)
        stale["source"]["stale"] = True
        stale["source"]["warning"] = (
            "Workbook reload failed. Showing the last successful snapshot."
        )
        self._failed_signature = signature
        self._stale_snapshot = stale
        return stale

    def _job_from_row(
        self,
        headers: tuple[Any, ...],
        row: tuple[Any, ...],
    ) -> dict[str, Any]:
        values = dict(zip(headers, row, strict=False))
        job = {
            key: self._json_value(values.get(header))
            for header, key in self.HEADER_KEYS.items()
        }
        job["job_url"] = self._safe_web_url(job["job_url"])
        job["application_url"] = self._safe_web_url(
            job["application_url"]
        )
        return job

    @staticmethod
    def _json_value(value: Any) -> Any:
        if value is None or isinstance(
            value,
            (str, int, float, bool),
        ):
            return value
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return str(value)

    @staticmethod
    def _safe_web_url(value: Any) -> str | None:
        if not isinstance(value, str) or len(value) > 4096:
            return None
        try:
            parsed = urlsplit(value)
        except ValueError:
            return None
        if (
            parsed.scheme not in {"http", "https"}
            or not parsed.hostname
            or parsed.username is not None
            or parsed.password is not None
        ):
            return None
        return value

    @staticmethod
    def _facet(
        jobs: list[dict[str, Any]],
        key: str,
    ) -> list[str]:
        return sorted({
            value
            for job in jobs
            if isinstance((value := job.get(key)), str)
            and value
        })

    @staticmethod
    def _metrics(jobs: list[dict[str, Any]]) -> dict[str, int]:
        def count_status(*statuses: str) -> int:
            return sum(
                job.get("status") in statuses
                for job in jobs
            )

        manual_queue = sum(
            job.get("company_rule") == "MANUAL"
            or job.get("application_method") in {"MANUAL", "REVIEW"}
            for job in jobs
        )

        return {
            "total_jobs": len(jobs),
            "applied": count_status("APPLIED"),
            "manual_queue": manual_queue,
            "needs_review": count_status("NEEDS_REVIEW"),
            "in_progress": count_status("OA", "INTERVIEW"),
            "offers": count_status("OFFER"),
            "rejected": count_status("REJECTED"),
        }
