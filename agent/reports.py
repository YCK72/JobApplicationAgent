from datetime import datetime
from zoneinfo import ZoneInfo
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import db
from .config import DATA, settings

COLUMNS = ['company', 'title', 'location', 'status', 'reason', 'resume', 'score', 'url', 'applied_at', 'updated', 'evidence', 'tailored_resume']


def export():
    day = datetime.now(ZoneInfo(settings().timezone)).date().isoformat()
    folder = DATA / 'reports'
    folder.mkdir(exist_ok=True)
    target = folder / f'applications-{day}.xlsx'
    book = Workbook()
    book.remove(book.active)
    rows = db.jobs()
    for name, statuses in [('Applied', {'applied'}), ('Needs Review', {'needs_review'}),
                           ('Pipeline', {'queued', 'running', 'skipped'})]:
        sheet = book.create_sheet(name)
        sheet.append([c.replace('_', ' ').title() for c in COLUMNS])
        for job in rows:
            if job['status'] not in statuses:
                continue
            sheet.append([job.get(k) or '' for k in COLUMNS])
            for cell in sheet[sheet.max_row]:
                # Force source-controlled text to remain text, never an Excel formula.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            sheet.row_dimensions[sheet.max_row].height = 45
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='164E45')
            cell.font = Font(color='FFFFFF', bold=True)
        for col, width in zip('ABCDEFGHIJKL', [24, 38, 28, 18, 64, 14, 10, 60, 28, 28, 40, 60]):
            sheet.column_dimensions[col].width = width
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        if sheet.max_row > 1:
            table = Table(displayName=name.replace(' ', ''), ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            sheet.add_table(table)
    temporary = target.with_name(f'.{target.stem}-{uuid4().hex}.tmp.xlsx')
    book.save(temporary)
    temporary.replace(target)
    return target
